import re, sys, html
sys.path.insert(0, '/sessions/youthful-gifted-turing/mnt/outputs/build')
from content import *
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GREEN = "1F5D3A"
NOTE = {  # fill, border, label color
    "tip": ("E8F3EC", "2E7D4F"),
    "remember": ("EAF0FA", "3A63B0"),
    "warn": ("FBE9E7", "C0392B"),
    "check": ("FFF4D6", "C98A00"),
    "diff": ("EFE9F7", "7B4FB0"),
}

def parts(text):
    out = []
    for tok in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text):
        if not tok: continue
        if tok.startswith("**"): out.append((tok[2:-2], True, False))
        elif tok.startswith("*") and len(tok) > 2: out.append((tok[1:-1], False, True))
        else: out.append((tok, False, False))
    return out

def add_runs(p, text, size=None, color=None):
    for seg, b, it in parts(text):
        r = p.add_run(seg)
        r.bold = b
        r.italic = it
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill)
    tcPr.append(sh)

def borders(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        v = kw.get(edge)
        e = OxmlElement(f'w:{edge}')
        if v:
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(v[0])); e.set(qn('w:color'), v[1])
        else:
            e.set(qn('w:val'), 'nil')
        b.append(e)
    tcPr.append(b)

def cell_margins(cell, top=50, bottom=50, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tcPr.append(m)

def set_width(table, widths):
    table.autofit = False
    for gc, w in zip(table._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w*1440)))
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = Inches(w)

def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:cantSplit'); trPr.append(e)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for m in ('left_margin', 'right_margin'): setattr(sec, m, Inches(1))
sec.top_margin = Inches(0.9); sec.bottom_margin = Inches(0.9)

st = doc.styles['Normal']
st.font.name = 'Calibri'; st.font.size = Pt(11.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.12

def style_heading(name, size, color, before, after):
    s = doc.styles[name]
    s.font.name = 'Georgia'; s.font.size = Pt(size); s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(color)
    s.element.rPr.rFonts.set(qn('w:ascii'), 'Georgia'); s.element.rPr.rFonts.set(qn('w:hAnsi'), 'Georgia')
    s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
style_heading('Heading 1', 24, GREEN, 0, 10)
style_heading('Heading 2', 14, "2E7D4F", 12, 4)

def bullet(text, style):
    p = doc.add_paragraph(style=style)
    add_runs(p, text)
    p.paragraph_format.space_after = Pt(3)
    return p

# ---- cover
for _ in range(5): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("♠ ♥ ♣ ♦"); r.font.size = Pt(28); r.font.color.rgb = RGBColor.from_string(GREEN)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(TITLE); r.font.name = 'Georgia'; r.font.size = Pt(44); r.bold = True
r.font.color.rgb = RGBColor.from_string(GREEN)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(SUBTITLE); r.font.size = Pt(16)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(TAGLINE); r.italic = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---- contents
p = doc.add_paragraph(); r = p.add_run("Contents"); r.font.name = 'Georgia'; r.font.size = Pt(24); r.bold = True
r.font.color.rgb = RGBColor.from_string(GREEN)
for title, _ in CH:
    p = doc.add_paragraph(title); p.paragraph_format.space_after = Pt(4)
    p.runs[0].font.size = Pt(13)

# ---- chapters
for ci, (title, blocks) in enumerate(CH):
    h = doc.add_heading(title, level=1)
    h.paragraph_format.page_break_before = True
    compact = title.startswith('12.')
    # rule under heading
    pPr = h._p.get_or_add_pPr(); pb = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    bt.set(qn('w:val'), 'single'); bt.set(qn('w:sz'), '10'); bt.set(qn('w:color'), GREEN); bt.set(qn('w:space'), '4')
    pb.append(bt); pPr.append(pb)
    for blk in blocks:
        k = blk[0]
        if k == 'h2':
            hh = doc.add_heading(blk[1], level=2)
            if compact: hh.paragraph_format.space_before = Pt(4)
        elif k == 'p':
            p = doc.add_paragraph(); add_runs(p, blk[1], size=(10 if compact else None))
        elif k == 'ul':
            for t in blk[1]: bullet(t, 'List Bullet')
        elif k == 'ol':
            # manual numbering to restart per list
            for i, t in enumerate(blk[1], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.35)
                p.paragraph_format.first_line_indent = Inches(-0.3)
                p.paragraph_format.space_after = Pt(3)
                rr = p.add_run(f"{i}.  "); rr.bold = True; rr.font.color.rgb = RGBColor.from_string(GREEN)
                add_runs(p, t)
        elif k == 'table':
            header, rows = blk[1], blk[2]
            n = len(header)
            total = 6.5
            if n == 2: widths = [2.9, 3.6] if header[0] in ("Word",) else [3.8, 2.7]
            elif n == 3: widths = [2.2, 2.0, 2.3]
            else: widths = [1.6, 2.3, 1.2, 1.4]
            if header[0] == "Word": widths = [1.9, 4.6]
            if header[0] == "Topic": widths = [1.4, 2.4, 2.7]
            if header[0] == "Hand": widths = [1.5, 4.3, 0.7]
            if len(header) == 4 and header[1] == "Shape": widths = [1.25, 1.25, 3.3, 0.7]
            t = doc.add_table(rows=1, cols=n); t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, htxt in zip(t.rows[0].cells, header):
                c.text = ""; pr = c.paragraphs[0]; rr = pr.add_run(htxt); rr.bold = True
                rr.font.color.rgb = RGBColor(255, 255, 255); rr.font.size = Pt(11)
                shade(c, GREEN); cell_margins(c, *( (25,25,120,120) if compact else (50,50,140,140) ))
                borders(c, bottom=(4, GREEN))
            for ri, row in enumerate(rows):
                cells = t.add_row().cells
                no_split(t.rows[-1])
                for c, txt in zip(cells, row):
                    c.text = ""; pr = c.paragraphs[0]; pr.paragraph_format.space_after = Pt(0)
                    add_runs(pr, txt, size=(9.5 if compact else 11))
                    cell_margins(c, *( (25,25,120,120) if compact else (50,50,140,140) ))
                    shade(c, "F3F7F4" if ri % 2 == 0 else "FFFFFF")
                    borders(c, bottom=(4, "CFDDD3"))
            for c in t.rows[0].cells: c.paragraphs[0].paragraph_format.space_after = Pt(0)
            set_width(t, widths)
            sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0); sp.paragraph_format.line_spacing = 0.4 if compact else 1.0
        elif k == 'note':
            fill, col = NOTE[blk[1]]
            t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
            c = t.rows[0].cells[0]; c.text = ""
            shade(c, fill); cell_margins(c, 100, 100, 180, 160)
            borders(c, left=(36, col))
            pr = c.paragraphs[0]; pr.paragraph_format.space_after = Pt(0)
            add_runs(pr, blk[2], size=11)
            no_split(t.rows[0]); set_width(t, [6.5])
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # remove trailing empty spacer so it cannot spill onto its own blank page
    lastp = doc.paragraphs[-1]
    if not lastp.text.strip() and lastp._element.getprevious() is not None and lastp._element.getprevious().tag.endswith('}tbl'):
        lastp._element.getparent().remove(lastp._element)

# footer page numbers
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
def fld(p, instr):
    r = p.add_run()
    for typ, txt in (('begin', None), (None, instr), ('end', None)):
        if typ:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), typ)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt
        r._r.append(e)
    r.font.size = Pt(9)
rr = fp.add_run(TITLE + "  |  Page "); rr.font.size = Pt(9); rr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
fld(fp, 'PAGE')

doc.save('/sessions/youthful-gifted-turing/mnt/outputs/Canasta_Made_Easy.docx')

# ================= HTML (online version) =================
def inl(t):
    t = html.escape(t)
    t = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", t)
    return re.sub(r"\*([^*]+?)\*", r"<em>\1</em>", t)

def slug(t): return re.sub(r'[^a-z0-9]+', '-', t.lower()).strip('-')

h = []
h.append(f"""<style>
:root {{ color-scheme: light }}
* {{ box-sizing: border-box }}
body {{ margin:0; font-family: Calibri, 'Segoe UI', Helvetica, Arial, sans-serif; color:#222; background:#f6f8f6; line-height:1.55; font-size:16px }}
header.hero {{ background:#1F5D3A; color:#fff; text-align:center; padding:48px 20px 40px }}
header.hero .suits {{ font-size:30px; letter-spacing:8px; opacity:.9 }}
header.hero h1 {{ font-family: Georgia, serif; font-size:44px; margin:8px 0 6px }}
header.hero p {{ margin:4px 0; font-size:17px }}
header.hero .tag {{ font-style:italic; opacity:.85; font-size:14px }}
.wrap {{ max-width:1000px; margin:0 auto; display:flex; gap:28px; padding:24px 16px }}
nav {{ width:250px; flex:none; position:sticky; top:12px; align-self:flex-start; background:#fff; border:1px solid #dbe5de; border-radius:10px; padding:14px 16px; font-size:14px; max-height:90vh; overflow:auto }}
nav h3 {{ margin:0 0 8px; font-family:Georgia,serif; color:#1F5D3A; font-size:16px }}
nav a {{ display:block; color:#1F5D3A; text-decoration:none; padding:4px 0; border-bottom:1px solid #eef3ef }}
nav a:hover {{ text-decoration:underline }}
main {{ flex:1; min-width:0 }}
section {{ background:#fff; border:1px solid #dbe5de; border-radius:12px; padding:24px 28px; margin-bottom:20px }}
section h2.ch {{ font-family:Georgia,serif; color:#1F5D3A; font-size:28px; margin:0 0 14px; padding-bottom:8px; border-bottom:3px solid #1F5D3A }}
h3 {{ font-family:Georgia,serif; color:#2E7D4F; font-size:19px; margin:20px 0 6px }}
p {{ margin:8px 0 }}
ul, ol {{ margin:8px 0 8px 22px; padding:0 }}
li {{ margin:4px 0 }}
table {{ border-collapse:collapse; width:100%; margin:12px 0; font-size:15px }}
th {{ background:#1F5D3A; color:#fff; text-align:left; padding:8px 12px }}
td {{ padding:8px 12px; border-bottom:1px solid #cfddd3 }}
tr:nth-child(even) td {{ background:#f3f7f4 }}
.note {{ border-left:6px solid; padding:10px 14px; margin:12px 0; border-radius:4px; font-size:15px }}
.note.tip {{ background:#E8F3EC; border-color:#2E7D4F }}
.note.remember {{ background:#EAF0FA; border-color:#3A63B0 }}
.note.warn {{ background:#FBE9E7; border-color:#C0392B }}
.note.check {{ background:#FFF4D6; border-color:#C98A00 }}
.note.diff {{ background:#EFE9F7; border-color:#7B4FB0 }}
@media (max-width:760px) {{ .wrap {{ flex-direction:column }} nav {{ width:auto; position:static }} header.hero h1 {{ font-size:32px }} section {{ padding:18px }} }}
@media print {{ nav {{ display:none }} body {{ background:#fff }} section {{ border:none; page-break-inside:auto }} header.hero {{ print-color-adjust:exact; -webkit-print-color-adjust:exact }} }}
</style>
<header class="hero"><div class="suits">&spades; &hearts; &clubs; &diams;</div><h1>{inl(TITLE)}</h1><p>{inl(SUBTITLE)}</p><p class="tag">{inl(TAGLINE)}</p></header>
<div class="wrap"><nav><h3>Contents</h3>""")
for title, _ in CH:
    h.append(f'<a href="#{slug(title)}">{html.escape(title)}</a>')
h.append("</nav><main>")
for title, blocks in CH:
    h.append(f'<section id="{slug(title)}"><h2 class="ch">{html.escape(title)}</h2>')
    for blk in blocks:
        k = blk[0]
        if k == 'h2': h.append(f"<h3>{html.escape(blk[1])}</h3>")
        elif k == 'p': h.append(f"<p>{inl(blk[1])}</p>")
        elif k == 'ul': h.append("<ul>" + "".join(f"<li>{inl(t)}</li>" for t in blk[1]) + "</ul>")
        elif k == 'ol': h.append("<ol>" + "".join(f"<li>{inl(t)}</li>" for t in blk[1]) + "</ol>")
        elif k == 'table':
            h.append("<table><thead><tr>" + "".join(f"<th>{html.escape(c)}</th>" for c in blk[1]) + "</tr></thead><tbody>")
            for row in blk[2]:
                h.append("<tr>" + "".join(f"<td>{inl(c)}</td>" for c in row) + "</tr>")
            h.append("</tbody></table>")
        elif k == 'note': h.append(f'<div class="note {blk[1]}">{inl(blk[2])}</div>')
    h.append("</section>")
h.append("</main></div>")
open('/sessions/youthful-gifted-turing/mnt/outputs/build/canasta_online.html', 'w').write("\n".join(h))
print("ok")
